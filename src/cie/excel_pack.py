"""
Excel deliverable — Retention investment scenario model.

WHY Excel: Finance and commercial managers want to change assumptions themselves
(margin, save rates, contact costs) and see the business case update. The data cannot
supply those numbers, so they are INPUTS, clearly marked, never hidden in code.

Conventions: blue font = hard-coded input; black = formula; yellow fill = assumption
the user should replace; every output is a live formula.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config as C

BLUE = Font(name="Arial", color="0000FF")
BLACK = Font(name="Arial")
BOLD = Font(name="Arial", bold=True)
HEAD = Font(name="Arial", bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
YELLOW = PatternFill("solid", fgColor="FFFF00")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
GBP = '£#,##0;(£#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'

# Illustrative defaults — ASSUMPTIONS, not findings (see README sheet)
DEFAULTS = {
    "1. Protect - Key Accounts":           (0.10, 40.0, "all"),
    "2. Rescue - High Value at Risk":      (0.25, 40.0, "all"),
    "3. Develop - Core Growth":            (0.10, 2.0, "all"),
    "4. Win Back - Core at Risk":          (0.15, 2.0, "all"),
    "5. Nurture - New Customers":          (0.10, 2.0, "all"),
    "6. Automate - Low Value Active":      (0.05, 0.5, "all"),
    "7. Deprioritise - Low Value Lapsing": (0.03, 0.5, "all"),
}


def _header(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")


def build(path=None):
    path = path or (C.ROOT / "excel" / "retention_investment_model.xlsx")
    seg = pd.read_csv(C.TABLES / "management_segments.csv")
    seg = seg[seg.management_segment.isin(DEFAULTS)]
    fy = pd.read_csv(C.TABLES / "sql_kpi_fy_comparison.csv")
    yoy = pd.read_csv(C.TABLES / "sql_kpi_yoy_retention.csv").iloc[0]
    hv = pd.read_csv(C.TABLES / "high_value_high_risk_top100.csv")

    wb = Workbook()
    # ------------------------------------------------------------------ README
    ws = wb.active
    ws.title = "README"
    notes = [
        ("Retention Investment Scenario Model", BOLD),
        ("Source: E-Commerce Customer Intelligence & LTV Engine (this repository). "
         "Data: UCI Online Retail II, 1 Dec 2009 – 9 Dec 2011.", BLACK),
        ("", BLACK),
        ("HOW TO USE", BOLD),
        ("1. Edit ONLY yellow cells (blue text) on the 'Assumptions' and 'Scenario' sheets.", BLACK),
        ("2. All other numbers are formulas and update automatically.", BLACK),
        ("", BLACK),
        ("WHAT IS CALCULATED vs ASSUMED", BOLD),
        ("CALCULATED from data: customers, trailing-12m revenue, revenue at risk "
         "(= trailing-12m revenue x modelled 12-month churn probability).", BLACK),
        ("ASSUMED (placeholders — replace with company figures): gross margin, save rate "
         "per segment, cost per contact. The dataset contains no cost, margin or campaign data.", BLACK),
        ("", BLACK),
        ("LIMITATIONS", BOLD),
        ("Revenue at risk is an exposure measure, not a loss forecast. It excludes spend "
         "decline by customers who stay (net revenue retention was 81%), so the true "
         "retention opportunity is larger.", BLACK),
        ("Customer metrics cover identified customers only (~85% of merchandise revenue).", BLACK),
        ("Save rates are hypotheses to be TESTED (hold-out control group), not facts.", BLACK),
    ]
    for i, (t, f) in enumerate(notes, 1):
        ws.cell(row=i, column=1, value=t).font = f
    ws.column_dimensions["A"].width = 120

    # ------------------------------------------------------------- Assumptions
    wa = wb.create_sheet("Assumptions")
    _header(wa, 1, ["Assumption", "Value", "Note"])
    wa["A2"], wa["B2"] = "Gross margin (%)", 0.40
    wa["C2"] = "ASSUMPTION — not in data. Replace with Finance's gross margin."
    wa["A3"], wa["B3"] = "Programme overhead (£, fixed)", 5000
    wa["C3"] = "ASSUMPTION — set-up cost (CRM configuration, analyst time)."
    for r, fmt in ((2, PCT), (3, GBP)):
        wa.cell(row=r, column=2).font = BLUE
        wa.cell(row=r, column=2).fill = YELLOW
        wa.cell(row=r, column=2).number_format = fmt
    for col, w in zip("ABC", (34, 14, 80)):
        wa.column_dimensions[col].width = w

    # ---------------------------------------------------------------- Scenario
    sc = wb.create_sheet("Scenario")
    cols = ["Management segment", "Customers", "Trailing-12m revenue (£)",
            "Revenue at risk (£)", "Save rate (assumption)", "Accounts contacted",
            "Cost per contact (£, assumption)", "Recovered revenue (£)",
            "Recovered gross profit (£)", "Programme cost (£)", "Net benefit (£)", "ROI (x)"]
    _header(sc, 1, cols)
    r0 = 2
    for i, row in enumerate(seg.itertuples(index=False)):
        r = r0 + i
        save, cost, _ = DEFAULTS[row.management_segment]
        sc.cell(row=r, column=1, value=row.management_segment).font = BLACK
        for col, val, fmt in ((2, int(row.customers), "#,##0"),
                              (3, round(float(row.revenue_t12m), 2), GBP),
                              (4, round(float(row.revenue_at_risk), 2), GBP)):
            c = sc.cell(row=r, column=col, value=val)
            c.font, c.number_format = BLUE, fmt
        for col, val, fmt in ((5, save, PCT), (7, cost, '£#,##0.00')):
            c = sc.cell(row=r, column=col, value=val)
            c.font, c.fill, c.number_format = BLUE, YELLOW, fmt
        c = sc.cell(row=r, column=6, value=f"=B{r}")
        c.number_format, c.fill = "#,##0", YELLOW
        sc.cell(row=r, column=8, value=f"=D{r}*E{r}").number_format = GBP
        sc.cell(row=r, column=9, value=f"=H{r}*Assumptions!$B$2").number_format = GBP
        sc.cell(row=r, column=10, value=f"=F{r}*G{r}").number_format = GBP
        sc.cell(row=r, column=11, value=f"=I{r}-J{r}").number_format = GBP
        sc.cell(row=r, column=12, value=f'=IF(J{r}=0,"n/a",K{r}/J{r})').number_format = '0.0"x"'
    rl = r0 + len(seg) - 1
    rt = rl + 1
    sc.cell(row=rt, column=1, value="Total (before overhead)").font = BOLD
    for col in (2, 3, 4, 6, 8, 9, 10, 11):
        L = get_column_letter(col)
        c = sc.cell(row=rt, column=col, value=f"=SUM({L}{r0}:{L}{rl})")
        c.font = BOLD
        c.number_format = "#,##0" if col in (2, 6) else GBP
    sc.cell(row=rt + 1, column=1, value="Less: programme overhead").font = BLACK
    sc.cell(row=rt + 1, column=11, value="=-Assumptions!B3").number_format = GBP
    sc.cell(row=rt + 2, column=1, value="Net benefit after overhead").font = BOLD
    c = sc.cell(row=rt + 2, column=11, value=f"=K{rt}+K{rt+1}")
    c.font, c.number_format = BOLD, GBP
    sc.cell(row=rt + 3, column=1, value="Break-even save rate, Rescue segment").font = BLACK
    rescue_row = r0 + list(seg.management_segment).index("2. Rescue - High Value at Risk")
    c = sc.cell(row=rt + 3, column=5,
                value=f"=IF(D{rescue_row}*Assumptions!$B$2=0,\"n/a\","
                      f"J{rescue_row}/(D{rescue_row}*Assumptions!$B$2))")
    c.number_format = PCT
    sc.cell(row=rt + 3, column=6,
            value="<- save rate at which Rescue outreach pays for itself").font = Font(
        name="Arial", italic=True)
    sc["E1"].comment = Comment("ASSUMPTION: share of probability-weighted revenue at risk "
                               "retained because of the intervention. Validate with a "
                               "randomised hold-out group.", "analyst")
    for j, w in enumerate((38, 11, 18, 16, 14, 13, 16, 16, 18, 15, 15, 9), 1):
        sc.column_dimensions[get_column_letter(j)].width = w
    sc.row_dimensions[1].height = 45
    sc.freeze_panes = "B2"

    # ---------------------------------------------------------------- KPI sheet
    k = wb.create_sheet("KPI_Summary")
    _header(k, 1, ["KPI (identified + guest unless stated)", "FY2010 (Dec09–Nov10)",
                   "FY2011 (Dec10–Nov11)", "Change"])
    f0, f1 = fy.iloc[0], fy.iloc[1]
    rows = [("Net merchandise revenue (£)", "net_merch_revenue", GBP),
            ("Identified-customer sales (£)", "identified_sales", GBP),
            ("Guest sales (£)", "guest_sales", GBP),
            ("Active identified customers", "active_customers", "#,##0"),
            ("Orders per customer", "orders_per_customer", "0.00"),
            ("AOV, identified (£)", "identified_aov", '£#,##0.00'),
            ("Return rate", "return_rate", PCT)]
    for i, (lab, col, fmt) in enumerate(rows, 2):
        k.cell(row=i, column=1, value=lab)
        for j, v in ((2, f0[col]), (3, f1[col])):
            c = k.cell(row=i, column=j, value=float(v))
            c.font, c.number_format = BLUE, fmt
        k.cell(row=i, column=4, value=f"=IF(B{i}=0,\"n/a\",C{i}/B{i}-1)").number_format = PCT
    n = len(rows) + 3
    k.cell(row=n, column=1, value="Customer retention FY10→FY11").font = BOLD
    c = k.cell(row=n, column=2, value=float(yoy.customer_retention_rate))
    c.font, c.number_format = BLUE, PCT
    k.cell(row=n + 1, column=1, value="Net revenue retention FY10→FY11").font = BOLD
    c = k.cell(row=n + 1, column=2, value=float(yoy.net_revenue_retention))
    c.font, c.number_format = BLUE, PCT
    k.cell(row=n + 3, column=1, value="Source: reports/tables/sql_kpi_fy_comparison.csv, "
                                      "sql_kpi_yoy_retention.csv (calculated from data)")
    for col, w in zip("ABCD", (42, 22, 22, 12)):
        k.column_dimensions[col].width = w

    # ------------------------------------------------------------- Rescue list
    rl_ws = wb.create_sheet("Rescue_Call_List")
    cols = ["priority_rank", "customer_id", "country", "revenue_t12m", "churn_prob_12m",
            "revenue_at_risk", "clv_12m", "recency_days", "avg_gap_days", "cadence_status",
            "at_risk_reason", "rfm_segment"]
    _header(rl_ws, 1, [c.replace("_", " ").title() for c in cols])
    for i, row in enumerate(hv[cols].itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            v = None if pd.isna(v) else (v.item() if hasattr(v, "item") else v)
            c = rl_ws.cell(row=i, column=j, value=v)
            c.font = BLUE
            if cols[j - 1] in ("revenue_t12m", "revenue_at_risk", "clv_12m"):
                c.number_format = GBP
            elif cols[j - 1] == "churn_prob_12m":
                c.number_format = PCT
            elif cols[j - 1] == "avg_gap_days":
                c.number_format = "0.0"
    rl_ws.freeze_panes = "C2"
    for j in range(1, len(cols) + 1):
        rl_ws.column_dimensions[get_column_letter(j)].width = 16
    rl_ws.column_dimensions["K"].width = 40

    for w in wb.worksheets:
        for row in w.iter_rows():
            for c in row:
                if c.font is None or c.font.name != "Arial":
                    c.font = Font(name="Arial", bold=c.font.bold if c.font else False,
                                  color=c.font.color if c.font else None,
                                  italic=c.font.italic if c.font else False)
    wb.save(path)
    print(f"  Excel model: {path.relative_to(C.ROOT)}")
    return path


if __name__ == "__main__":
    build()
